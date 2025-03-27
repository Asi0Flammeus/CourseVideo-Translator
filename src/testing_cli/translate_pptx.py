import os
import re
import sys
import argparse
from typing import Optional, Dict
from pathlib import Path
from functools import lru_cache
from tqdm import tqdm
from pptx import Presentation
import deepl
from dotenv import load_dotenv

# Load .env from parent directory
dotenv_path = Path(__file__).parent.parent / '.env'
load_dotenv(dotenv_path)
# Language mappings
SUPPORTED_LANGUAGES: Dict[str, str] = {
    "cs": "Czech",
    "de": "German",
    "en": "English",
    "es": "Spanish",
    "et": "Estonian",
    "fi": "Finnish",
    "fr": "French",
    "id": "Indonesian",
    "it": "Italian",
    "ja": "Japanese",
    "nb-NO": "Norwegian (Bokmål)",
    "pt": "Portuguese",
    "ru": "Russian",
    "vi": "Vietnamese",
    "zh-Hans": "Chinese (Simplified)"
}

DEEPL_LANGUAGE_MAPPING: Dict[str, str] = {
    "cs": "CS",
    "de": "DE",
    "en": "EN-US",
    "es": "ES",
    "et": "ET",
    "fi": "FI",
    "fr": "FR",
    "id": "ID",
    "it": "IT",
    "ja": "JA",
    "nb-NO": "NB",
    "pt": "PT",
    "ru": "RU",
    "vi": "VI",
    "zh-Hans": "ZH"
}

class TranslationError(Exception):
    """Custom exception for translation errors"""
    pass

class PPTXTranslator:
    def __init__(self, auth_key: str = None):
        """Initialize translator with DeepL API key"""
        self.auth_key = auth_key or os.getenv('DEEPL_API_KEY')
        if not self.auth_key:
            raise ValueError("DeepL API key is required. Set DEEPL_API_KEY environment variable or pass it directly.")
        
        self.translator = deepl.Translator(self.auth_key)
        self.translation_cache = {}

    @lru_cache(maxsize=1000)
    def get_translation(self, text: str, target_language: str) -> str:
        """
        Translate text using DeepL API with caching
        
        Args:
            text: Text to translate
            target_language: Target language code
            
        Returns:
            Translated text
            
        Raises:
            TranslationError: If translation fails
        """
        if not text.strip():
            return text
        
        cache_key = (text, target_language)
        if cache_key in self.translation_cache:
            return self.translation_cache[cache_key]
        
        try:
            deepl_target_lang = DEEPL_LANGUAGE_MAPPING.get(target_language)
            if not deepl_target_lang:
                raise TranslationError(f"Unsupported target language: {target_language}")
                
            result = self.translator.translate_text(text, target_lang=deepl_target_lang)
            translated_text = result.text
            self.translation_cache[cache_key] = translated_text
            return translated_text
        except Exception as e:
            raise TranslationError(f"Translation failed: {str(e)}")

    def is_exception_text(self, text: str, source_lang: str, target_lang: str) -> Optional[str]:
        """
        Handle special text patterns that shouldn't be translated
        
        Args:
            text: Text to check
            source_lang: Source language code
            target_lang: Target language code
            
        Returns:
            Modified text if it's an exception, None otherwise
        """
        # Language code exception (e.g., "- EN" -> "- FR")
        if f"- {source_lang.upper()}" in text:
            return text.replace(f"- {source_lang.upper()}", f"- {target_lang.upper()}")
        return None

    def translate_pptx(self, source_lang: str, target_lang: str, input_path: str, 
                      output_path: str, use_exceptions: bool = True) -> None:
        """
        Translate PowerPoint presentation
        
        Args:
            source_lang: Source language code
            target_lang: Target language code
            input_path: Path to input PPTX file
            output_path: Path to save translated PPTX
            use_exceptions: Whether to check for exception patterns
            
        Raises:
            TranslationError: If translation fails
        """
        try:
            # Validate languages
            if source_lang not in SUPPORTED_LANGUAGES:
                raise TranslationError(f"Unsupported source language: {source_lang}")
            if target_lang not in SUPPORTED_LANGUAGES:
                raise TranslationError(f"Unsupported target language: {target_lang}")

            # Load presentation
            prs = Presentation(input_path)
            
            # Count total translation units for progress bar
            total_runs = sum(
                len(shape.text_frame.paragraphs)
                for slide in prs.slides
                for shape in slide.shapes
                if shape.has_text_frame
            )

            # Translate presentation
            with tqdm(total=total_runs, desc="Translating slides") as pbar:
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if not shape.has_text_frame:
                            continue
                        
                        for paragraph in shape.text_frame.paragraphs:
                            for run in paragraph.runs:
                                if not run.text.strip():
                                    continue
                                    
                                if use_exceptions:
                                    exception_result = self.is_exception_text(
                                        run.text, source_lang, target_lang
                                    )
                                    if exception_result:
                                        run.text = exception_result
                                        continue
                                
                                run.text = self.get_translation(run.text, target_lang)
                            pbar.update(1)

            # Ensure output directory exists
            output_dir = Path(output_path).parent
            output_dir.mkdir(parents=True, exist_ok=True)
            
            # Save translated presentation
            prs.save(output_path)
            
        except Exception as e:
            raise TranslationError(f"Failed to translate presentation: {str(e)}")


def list_input_files() -> list:
    """List all PPTX files in the input directory"""
    input_dir = Path("../inputs/pptx")
    return list(input_dir.glob("*.pptx"))

def prompt_language_selection(prompt_text: str) -> str:
    """
    Interactive language selection prompt
    
    Args:
        prompt_text: Text to display before language options
        
    Returns:
        Selected language code
    """
    print(f"\n{prompt_text}")
    for code, name in sorted(SUPPORTED_LANGUAGES.items()):
        print(f"{code}: {name}")
    
    while True:
        lang = input("\nEnter language code: ").lower()
        if lang in SUPPORTED_LANGUAGES:
            return lang
        print("Invalid language code. Please try again.")


def main():
    """CLI entry point"""
    # Set the base path
    base_path = "../../../vault/DBxPBN/workspace/assets/educational-content/courses/Free/"
    
    try:
        # Get and sort the folders in the base path
        folders = [f for f in os.listdir(base_path) if os.path.isdir(os.path.join(base_path, f))]
        folders.sort()
        
        # Print available folders
        print("\n=== PPTX Translator ===")
        print("\nAvailable folders:")
        for i, folder in enumerate(folders, 1):
            print(f"{i}. {folder}")
            
        # Get user input for folder selection
        while True:
            try:
                selection = int(input("\nEnter the number of the folder to process (1-{}): ".format(len(folders))))
                if 1 <= selection <= len(folders):
                    break
                print("Invalid selection. Please try again.")
            except ValueError:
                print("Please enter a valid number.")
        
        selected_folder = folders[selection-1]
        
        # Language selection
        source_lang = prompt_language_selection("Select source language:")
        target_lang = prompt_language_selection("Select target language:")
        
        # Define source and target folder paths
        source_folder_path = os.path.join(base_path, selected_folder, f"translation/{source_lang}/v001")
        target_folder_path = os.path.join(base_path, selected_folder, f"translation/{target_lang}/v001")
        
        # Check if source folder exists
        if not os.path.exists(source_folder_path):
            raise Exception(f"Source folder does not exist: {source_folder_path}")
        
        # Create target folder if it doesn't exist
        os.makedirs(target_folder_path, exist_ok=True)
        
        # Show confirmation
        print("\nPlease confirm the following paths:")
        print(f"\nSource folder: {source_folder_path}")
        print(f"Target folder: {target_folder_path}")
        
        confirm = input("\nProceed with translation? (y/n): ").lower()
        if confirm != 'y':
            print("Translation cancelled.")
            return
        
        # Get all subfolders from source
        subfolders = [f for f in os.listdir(source_folder_path) if os.path.isdir(os.path.join(source_folder_path, f))]
        subfolders.sort()
        
        # Initialize translator
        translator = PPTXTranslator()
        
        # Process each subfolder
        for subfolder in subfolders:
            source_subfolder_path = os.path.join(source_folder_path, subfolder)
            target_subfolder_path = os.path.join(target_folder_path, subfolder)
            
            # Create target subfolder if it doesn't exist
            os.makedirs(target_subfolder_path, exist_ok=True)
            
            # Get all PPTX files in the source subfolder
            pptx_files = [f for f in os.listdir(source_subfolder_path) if f.lower().endswith('.pptx')]
            
            if not pptx_files:
                print(f"\nNo PPTX files found in {subfolder}")
                continue
                
            print(f"\nProcessing {subfolder}...")
            
            # Process each PPTX file
            for pptx_file in pptx_files:
                print(f"\nTranslating: {pptx_file}")
                
                # Construct input and output paths
                input_path = os.path.join(source_subfolder_path, pptx_file)
                output_path = os.path.join(target_subfolder_path, pptx_file)
                
                # Check if translated file already exists
                if os.path.exists(output_path):
                    print(f"Skipping {pptx_file} - translation already exists at: {output_path}")
                    continue
                
                try:
                    translator.translate_pptx(
                        source_lang=source_lang,
                        target_lang=target_lang,
                        input_path=input_path,
                        output_path=output_path
                    )
                    print(f"Translation completed for {pptx_file}")
                    print(f"Saved to: {output_path}")
                except Exception as e:
                    print(f"Error translating {pptx_file}: {str(e)}")
                    continue
                    
    except TranslationError as e:
        print(f"\nTranslation error: {str(e)}")
        sys.exit(1)
    except Exception as e:
        print(f"\nUnexpected error: {str(e)}")
        sys.exit(1)

if __name__ == "__main__":
    main()

